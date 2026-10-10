"""Apply consistent report layout using the image's existing python-docx tooling."""
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

document = Document(sys.argv[1])
normal = document.styles['Normal']
normal.font.name = 'Liberation Sans'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
for section in document.sections:
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Page ')
    for instruction, following in [('PAGE', ' of '), ('NUMPAGES', '')]:
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), instruction)
        footer._p.append(field)
        footer.add_run(following)
document.save(sys.argv[1])

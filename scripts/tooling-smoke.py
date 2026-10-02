"""Exercise image/document generation and headless conversion in the built image."""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation


def main(output: Path) -> None:
    required = ["uv", "rg", "jq", "zip", "unzip", "pandoc", "libreoffice", "pdftotext"]
    missing = [name for name in required if not shutil.which(name)]
    if missing:
        raise RuntimeError(f"Missing image tools: {', '.join(missing)}")
    output.mkdir(parents=True, exist_ok=False)
    data = pd.DataFrame({"day": [1, 2, 3], "value": [2, 4, 3]})
    data.to_excel(output / "report.xlsx", index=False)
    assert load_workbook(output / "report.xlsx").active["B3"].value == 4
    plt.plot(data.day, data.value)
    plt.savefig(output / "chart.png")
    plt.close()
    with Image.open(output / "chart.png") as image:
        image.verify()
    document = Document()
    document.add_heading("Assistant tooling smoke test", 0)
    document.add_paragraph("Generated locally with the Python document tools.")
    document.save(output / "report.docx")
    assert Document(output / "report.docx").paragraphs[0].text == "Assistant tooling smoke test"
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[0]).shapes.title.text = "Assistant tooling"
    slides.save(output / "report.pptx")
    assert len(Presentation(output / "report.pptx").slides) == 1
    subprocess.run(["pandoc", "-f", "markdown", "-t", "html", "-o", str(output / "report.html")], input="# Assistant tooling\n", text=True, check=True, timeout=30)
    with tempfile.TemporaryDirectory(prefix="assistant-lo-") as profile:
        subprocess.run(["libreoffice", f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(output), str(output / "report.docx")], check=True, capture_output=True, timeout=90)
    extracted = subprocess.run(["pdftotext", str(output / "report.pdf"), "-"], check=True, capture_output=True, text=True, timeout=30)
    assert "Assistant tooling smoke test" in extracted.stdout
    print(json.dumps({"generated": sorted(p.name for p in output.iterdir()), "status": "passed"}))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: tooling-smoke.py NEW_OUTPUT_DIRECTORY")
    main(Path(sys.argv[1]).resolve())
